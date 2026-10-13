from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\raulm\Downloads\Bakcend-KT")
SOURCE = ROOT / "docs" / "CLIENT-FEATURE-GUIDE.md"
OUTPUT = ROOT / "docs" / "CLIENT-FEATURE-GUIDE.docx"

SCREENSHOTS = [
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-add8f94b-8816-4f09-8a28-74a4af167c46.png"), "Fiscal Year selector dan status periode"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-838ae8f5-7679-4098-85ae-32b5104a0f36.png"), "Dropdown mata uang pada Donor dan Grant"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-1abef574-1ace-4abb-a3de-b369faa31a56.png"), "Dropdown bank account pada Donor dan Grant"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-1c9bb8f1-acc7-40c0-9ff6-2b49edac4688.png"), "Dropdown Budget Categories"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-3412ef5b-5032-4115-845a-3753fc756410.png"), "Halaman Programs and Projects"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-f777ec21-5716-4297-b543-b2964992db38.png"), "Project Timeline dan Workplan dengan Fiscal Year"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-3a2ccf8c-ad41-4f0b-b40c-1eeda6e7c8ac.png"), "Master Programs dengan pilihan tahun fiskal"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-6ddd218d-7039-492e-800d-5a8def17fc8f.png"), "Financial report dan General Ledger view"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-b6ed1bbb-c5f9-4ca2-bfaf-1d4fc0200672.png"), "Donor and Grant report"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-d7d9784d-fb61-4066-946e-a0f6bffeecc1.png"), "Cash Advance dan expense monitoring"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-29f8e776-3c53-44eb-bfb2-0978e15d366b.png"), "Reimbursement dan action menu"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-90e3924e-6f48-4a7c-9181-b7b85f1914bf.png"), "Tax PPh dan action column"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-36d25f87-3f31-47aa-a444-533fdf1db6d8.png"), "VAT transaction dan action menu"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-5d677f23-bcb0-4ffb-8ad5-da24da987ab9.png"), "Procurement RFQ action menu"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-4eaf48fe-eb42-4a8a-9046-dfbe81774e22.png"), "Bank reconciliation status dan action"),
    (Path(r"C:\Users\raulm\AppData\Local\Temp\codex-clipboard-55f9f0d6-759a-4e30-bf8b-b4f3426aa3f7.png"), "Chart of Accounts parent account dropdown"),
]

SCREENSHOT_BY_HEADING = {
    "1.1 Fiscal Year": SCREENSHOTS[0],
    "3.1 Finance & Accounting": SCREENSHOTS[1],
    "Bank Accounts": SCREENSHOTS[2],
    "Currencies": SCREENSHOTS[1],
    "3.3 Budget & Reporting": SCREENSHOTS[3],
    "Programs": SCREENSHOTS[4],
    "Project Timeline / Workplan": SCREENSHOTS[5],
    "Financial Reports": SCREENSHOTS[7],
    "Donor / Grant Reports": SCREENSHOTS[8],
    "Cash Advances": SCREENSHOTS[9],
    "Reimbursements": SCREENSHOTS[10],
    "PPh dan VAT/PPN": SCREENSHOTS[11],
    "VAT/PPN": SCREENSHOTS[12],
    "RFQ & CBA": SCREENSHOTS[13],
    "Bank Reconciliation": SCREENSHOTS[14],
    "Chart of Accounts": SCREENSHOTS[15],
}

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F6F8"
BORDER = "D9D9D9"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def cell_borders(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text.strip())
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, rows):
    parsed = []
    for row in rows:
        values = [x.strip() for x in row.strip().strip("|").split("|")]
        parsed.append(values)
    if len(parsed) < 2:
        return
    header = parsed[0]
    body = [r for r in parsed[2:] if not all(re.fullmatch(r":?-+:?", x) for x in r)]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, value in enumerate(header):
        cell = table.rows[0].cells[index]
        shade(cell, BLUE)
        cell_borders(cell)
        set_cell_text(cell, value, bold=True, color="FFFFFF", size=8.5)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_index, row in enumerate(body):
        cells = table.add_row().cells
        for index in range(len(header)):
            value = row[index] if index < len(row) else ""
            cell_borders(cells[index])
            if row_index % 2 == 1:
                shade(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], value, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_inline_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    # Basic bold handling for markdown emphasis.
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = p.add_run(value)
        run.bold = bold
        run.font.name = "Aptos"
        run.font.size = Pt(10)
    return p

def add_feature_screenshot(doc, screenshot):
    path, caption = screenshot
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(5.9))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", caption)
    doc_pr.set("title", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run("Bukti tampilan: " + caption)
    cr.italic = True
    cr.font.name = "Aptos"
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor.from_string("666666")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    # Header and footer for client-facing handover.
    header = section.header.paragraphs[0]
    header.text = "Eco Finance Kaoem Telapak | Feature Guide and SOP"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.text = "Client Handover Document"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.name = "Aptos"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("666666")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    table_buffer = []
    in_code = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(stripped)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue
        if stripped.startswith("|"):
            table_buffer.append(line)
            continue
        if table_buffer:
            add_table(doc, table_buffer)
            table_buffer = []
        if not stripped or stripped == "---":
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.add_run(stripped[2:].strip())
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1)) - 1
            heading_text = heading.group(2).strip()
            p = doc.add_heading(heading_text, level=level)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            screenshot = SCREENSHOT_BY_HEADING.get(heading_text)
            if screenshot:
                add_feature_screenshot(doc, screenshot)
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            r = p.add_run(stripped[2:])
            r.italic = True
            r.font.name = "Aptos"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor.from_string("555555")
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            p.paragraph_format.space_after = Pt(2)
            p.add_run((bullet or numbered).group(1))
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(10)
            continue
        add_inline_paragraph(doc, stripped)

    if table_buffer:
        add_table(doc, table_buffer)
    add_inline_paragraph(doc, "Catatan bukti tampilan: screenshot ditempatkan langsung pada bagian fitur terkait. Screenshot untuk halaman yang belum tersedia memerlukan akses staging yang sudah login agar bukti yang digunakan berasal dari environment client yang benar.")
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__":
    main()
